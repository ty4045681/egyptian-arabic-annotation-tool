"""Training tar blocks open rounds; Excel/JSON keep blocked audit rows."""

from __future__ import annotations

import json
import uuid
from pathlib import Path

import pytest
from openpyxl import load_workbook

import db
from annotation_metadata.export_metadata import EXCEL_HEADERS
from annotation_metadata.postgres_backup import (
    provenance_counts, provenance_relationships,
)
from annotation_quality.contracts import CrossCheckDecisionCommand
from annotation_quality.service import decide_cross_check
from export import export_xlsx
from manage_state import export_json
from scripts import export_top_annotators_tar as exporter
from tests.test_cross_check_admin import (
    SECRET_A,
    SECRET_B,
    _admin_session,
    queue_rounds,
)
from tests.test_cross_check_claim import enable_cross_check
from tests.test_cross_check_submit import (
    IDENTICAL,
    bob_claim,
    complete,
    text_segments,
)
from tests.test_cross_check_stats import _seed_annotated_duration
from tests.test_api import login


def _snapshot(database, **kwargs):
    return exporter.load_snapshot(
        database,
        kwargs.pop("top_n", None),
        kwargs.pop("min_seconds", 0.0),
        **kwargs,
    )


def _published_version(task_id):
    with db.db_conn() as conn:
        return str(conn.execute(
            """SELECT current_published_version_id
               FROM annotation_tasks WHERE id = %s""",
            (task_id,),
        ).fetchone()[0])


def _round_versions(round_id):
    with db.db_conn() as conn:
        return conn.execute(
            """SELECT original_version_id, secondary_version_id, revision,
                      secondary_annotator_id
               FROM cross_check_rounds WHERE id = %s""",
            (round_id,),
        ).fetchone()


def test_open_rounds_excluded_from_training_snapshot(client, seed_tasks, database):
    task_ids = _seed_annotated_duration(client, seed_tasks, 2, 60.0)
    enable_cross_check(enabled=True, sampling_rate_bps=10000)
    login(client, "bob")
    in_progress = bob_claim(client)
    blocked_id = in_progress["task_id"]
    remaining = next(task_id for task_id in task_ids if task_id != blocked_id)

    snapshot = _snapshot(database)
    exported = {item.task_id for item in snapshot.segments}
    assert str(blocked_id) not in exported
    assert str(remaining) in exported
    assert snapshot.snapshot_at is not None
    assert snapshot.excluded_open_quality_tasks == 1
    assert snapshot.excluded_open_quality_seconds == pytest.approx(60.0)

    queued, _ = complete(
        client, in_progress,
        segments=text_segments(in_progress, "totally different transcript here"),
    )
    assert queued.status_code == 200, queued.json
    assert queued.json["cross_check"]["state"] == "awaiting_review"
    snapshot = _snapshot(database)
    exported = {item.task_id for item in snapshot.segments}
    assert str(blocked_id) not in exported
    assert str(remaining) in exported
    assert snapshot.excluded_open_quality_tasks == 1


def test_passed_export_only_final_version(client, seed_tasks, database):
    _seed_annotated_duration(client, seed_tasks, 1, 60.0, text=IDENTICAL)
    enable_cross_check(enabled=True, sampling_rate_bps=10000)
    login(client, "bob")
    assignment = bob_claim(client)
    passed, _ = complete(
        client, assignment, segments=text_segments(assignment, IDENTICAL),
    )
    assert passed.json["cross_check"]["state"] == "passed"
    task_id = assignment["task_id"]
    published = _published_version(task_id)
    snapshot = _snapshot(database)
    assert {item.task_id for item in snapshot.segments} == {str(task_id)}
    assert all(item.text for item in snapshot.segments)
    quality = {item["task_id"]: item for item in snapshot.exported_task_quality}
    assert quality[str(task_id)]["final_version_id"] == published
    assert quality[str(task_id)]["round_state"] == "passed"


def test_adjudicated_export_only_final_version(client, seed_tasks, database):
    queued = queue_rounds(client, seed_tasks, 1)
    round_id = queued[0]["round_id"]
    original_id, secondary_id, revision, _ = _round_versions(round_id)
    command = CrossCheckDecisionCommand.model_validate({
        "operation_id": str(uuid.uuid4()),
        "expected_revision": int(revision),
        "expected_original_version_id": str(original_id),
        "expected_secondary_version_id": str(secondary_id),
        "decision": "secondary",
        "reason": "export only B",
    })
    result = decide_cross_check(str(_admin_session()["id"]), str(round_id), command)
    assert result["final_version_id"] == str(secondary_id)
    snapshot = _snapshot(database)
    exported_ids = {item.task_id for item in snapshot.segments}
    assert str(queued[0]["task_id"]) in exported_ids
    quality = {item["task_id"]: item for item in snapshot.exported_task_quality}
    assert quality[str(queued[0]["task_id"])]["final_version_id"] == str(secondary_id)
    assert quality[str(queued[0]["task_id"])]["round_state"] == "adjudicated"
    assert quality[str(queued[0]["task_id"])]["outcome"] == "secondary"


def test_admin_edited_null_submitter_is_exported(client, seed_tasks, database):
    queued = queue_rounds(
        client, seed_tasks, 1,
        original_text=SECRET_A, secondary_text=SECRET_B,
    )
    round_id = queued[0]["round_id"]
    detail_row = _round_versions(round_id)
    original_id, secondary_id, revision, _ = detail_row
    with db.db_conn() as conn:
        original_segments = [
            {
                "id": row[0], "start": row[1], "end": row[2],
                "duration": row[3], "text": "admin edited " + (row[4] or ""),
                "exclude_from_training": row[5],
            }
            for row in conn.execute(
                """SELECT segment_id, start_s, end_s, duration, text,
                          exclude_from_training
                   FROM segments WHERE version_id = %s ORDER BY segment_id""",
                (original_id,),
            ).fetchall()
        ]
    command = CrossCheckDecisionCommand.model_validate({
        "operation_id": str(uuid.uuid4()),
        "expected_revision": int(revision),
        "expected_original_version_id": str(original_id),
        "expected_secondary_version_id": str(secondary_id),
        "decision": "edited",
        "reason": "admin rewrite for export",
        "base": "original",
        "segments": original_segments,
        "target_status": "annotated",
    })
    result = decide_cross_check(str(_admin_session()["id"]), str(round_id), command)
    final_id = result["final_version_id"]
    with db.db_conn() as conn:
        submitter, credited = conn.execute(
            """SELECT submitted_by_user_id, credited_annotator_id
               FROM annotation_versions WHERE id = %s""",
            (final_id,),
        ).fetchone()
    assert submitter is None
    assert credited is not None
    snapshot = _snapshot(database)
    exported = {item.task_id for item in snapshot.segments}
    assert str(queued[0]["task_id"]) in exported
    assert any("admin edited" in (item.text or "") for item in snapshot.segments)


def test_excel_one_row_per_task_including_blocked(client, seed_tasks, tmp_path):
    queued = queue_rounds(client, seed_tasks, 1)
    ready_ids = seed_tasks(1, folder="ready")
    from tests.test_scene_claims import source
    source(ready_ids[0], "airport", "high", batch="excel-ready-batch")
    login(client, "alice")
    claimed = client.post(
        "/api/assignment/claim", json={"source_scene": "airport"},
    )
    assert claimed.status_code == 200, claimed.json
    done, _ = complete(
        client, claimed.json, segments=text_segments(claimed.json, IDENTICAL),
    )
    assert done.status_code == 200, done.json
    output = tmp_path / "audit.xlsx"
    count = export_xlsx(output)
    assert count >= 2
    rows = list(load_workbook(output, read_only=True).active.iter_rows(values_only=True))
    assert list(rows[0]) == EXCEL_HEADERS
    quality_idx = EXCEL_HEADERS.index("Quality state")
    eligible_idx = EXCEL_HEADERS.index("Training eligible")
    body = rows[1:]
    states = {row[quality_idx] for row in body}
    assert "awaiting_review" in states
    assert any(row[eligible_idx] == "no" for row in body)
    assert any(row[eligible_idx] == "yes" for row in body)
    keys = [(row[1], row[2]) for row in body]
    assert len(keys) == len(set(keys))
    assert queued[0]["task_id"]


def test_export_json_keeps_pending_review_and_adds_eligibility(
        client, seed_tasks, tmp_path):
    queued = queue_rounds(client, seed_tasks, 1)
    out = tmp_path / "legacy-json"
    result = export_json(out)
    assert result["count"] >= 1
    payloads = [
        json.loads(path.read_text())
        for path in Path(out).rglob("*.json")
        if path.name not in {"assignments.json", "export_manifest.json"}
    ]
    blocked = [
        item for item in payloads
        if item.get("quality_state") == "awaiting_review"
    ]
    assert blocked
    assert all(item["training_eligible"] is False for item in blocked)
    assert all("segments" in item for item in payloads)
    assert queued[0]["task_id"]


def test_backup_checklist_includes_quality_tables(database):
    with db.db_conn() as conn:
        counts = provenance_counts(conn)
        relations = provenance_relationships(conn)
    assert "cross_check_settings" in counts
    assert "cross_check_rounds" in counts
    assert "task_annotation_participants" in counts
    assert counts["cross_check_settings"] == 1
    assert relations["ok"]
    assert "rounds_without_task" in relations["dangling"]
    assert "in_progress_rounds_without_assignment" in relations["dangling"]
