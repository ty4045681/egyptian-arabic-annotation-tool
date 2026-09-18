"""Excel provenance columns and real training-exporter sidecar CLI."""

from __future__ import annotations

import json
import os
import subprocess
import sys
import tarfile
from pathlib import Path

import numpy as np
import soundfile as sf
from openpyxl import load_workbook

import annotation_repository as repo
import db
from annotation_metadata.contracts import TaskFilter
from annotation_metadata.export_metadata import AUDIO_LEVEL_NOTICE, EXCEL_HEADERS
from export import export_xlsx
from scripts import export_top_annotators_tar as exporter
from tests.provenance_fixtures import build_rich_provenance
from tests.test_repository import full_segments, make_user


def test_excel_includes_source_and_human_columns(database, seed_tasks, tmp_path):
    fixture = build_rich_provenance(seed_tasks)
    output = tmp_path / "annotations.xlsx"
    count = export_xlsx(output)
    assert count >= 2
    workbook = load_workbook(output, read_only=True)
    rows = list(workbook.active.iter_rows(values_only=True))
    assert list(rows[0]) == EXCEL_HEADERS
    body = rows[1:]
    assert any(row[5] and "airport" in str(row[5]) for row in body)
    assert any(row[7] and "airport:high" in str(row[7]) for row in body)
    assert any(row[9] == "mixed" for row in body)
    assert any(row[10] and "airport" in str(row[10]) and "shopping" in str(row[10]) for row in body)
    filtered = tmp_path / "airport.xlsx"
    assert export_xlsx(filtered, filters=TaskFilter(source_scene="airport")) >= 1
    filtered_rows = list(load_workbook(filtered, read_only=True).active.iter_rows(values_only=True))
    assert all(
        "shopping" not in str(row[5] or "") or "airport" in str(row[5] or "")
        for row in filtered_rows[1:]
    )
    assert fixture["rich_id"]  # rich task is in the unfiltered workbook corpus


def test_training_exporter_cli_writes_audio_level_sidecar(database, seed_tasks, tmp_path):
    task_id = seed_tasks(1, duration=12.0)[0]
    user, _ = make_user("trainer")
    assignment = repo.claim(user["fence"])
    assert assignment["task_id"] == task_id
    repo.complete(
        user["fence"], assignment["lease_token"], 0, "annotated", [],
        full_segments(assignment, "train"), str(__import__("uuid").uuid4()), "train-complete",
        scene_review={"status": "confirmed", "scene_codes": ["restaurant"]},
    )
    with db.db_conn() as conn, conn.cursor() as cur:
        from annotation_metadata.repository import ensure_batch, sync_source_record
        from annotation_metadata.contracts import NormalizedSource
        batch = ensure_batch(cur, "acceptance-train-batch")
        sync_source_record(
            cur, task_id=task_id, batch_id=batch,
            source=NormalizedSource(
                record_key="train-key", scene_code="restaurant", confidence="high",
                confidence_basis="training fixture", source_url="https://example.com/train",
                video_id="trainvid", provider="youtube",
                raw_record={"video_id": "trainvid"}, content_digest="ff" * 32,
            ),
        )
    audio_dir = tmp_path / "audio-root"
    wav = audio_dir / "audio-000.wav"
    wav.parent.mkdir(parents=True)
    sf.write(wav, np.zeros(12 * exporter.TARGET_SAMPLE_RATE, dtype=np.float32),
             exporter.TARGET_SAMPLE_RATE, subtype="PCM_16")
    output = tmp_path / "dataset.tar"
    env = {**os.environ, "ANNOTATION_DB_DSN": database, "ANNOTATION_BACKUP_DSN": database}
    subprocess.run(
        [
            sys.executable, "scripts/export_top_annotators_tar.py",
            "--output", str(output), "--audio-dir", str(audio_dir),
            "--min-hours", "0", "--top-n", "1", "--dsn-env", "ANNOTATION_BACKUP_DSN",
        ],
        cwd=Path(__file__).resolve().parents[1],
        check=True, env=env,
    )
    with tarfile.open(output, "r") as archive:
        members = {member.name for member in archive.getmembers()}
        items = json.load(archive.extractfile("data.json"))
        sidecar = json.load(archive.extractfile("scene_metadata.json"))
        selection = json.load(archive.extractfile("export_metadata.json"))
    assert "data.json" in members
    assert "scene_metadata.json" in members
    assert set(items[0]) <= {"audio", "text", "asr_text", "annotator",
                             "annotation_started_at", "annotation_completed_at",
                             "annotation_elapsed_seconds"}
    assert "scene_code" not in items[0]
    assert sidecar["level"] == "audio"
    assert AUDIO_LEVEL_NOTICE in sidecar["notice"]
    sidecar_audio = [item["audio"] for item in sidecar["segments"]]
    assert sidecar_audio == [item["audio"] for item in items]
    assert {item["task_id"] for item in sidecar["segments"]} == {
        task["task_id"] for task in sidecar["tasks"]
    }
    assert sidecar["tasks"][0]["task_id"] == str(task_id)
    assert any(source["scene_code"] == "restaurant" for source in sidecar["tasks"][0]["sources"])
    assert sidecar["tasks"][0]["reviews"]
    assert selection["scene_metadata"] == "scene_metadata.json"
    assert selection["applied_task_filters"]["mode"] == "all_data"
    assert selection["data_json_fields"] == ["audio", "text", "asr_text"]
    for item in items:
        assert item["audio"] in members
